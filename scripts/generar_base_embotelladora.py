#!/usr/bin/env python3
"""
Convierte el maestro sintetico de embotelladora en un Excel presentable, del
tipo que se muestra en una reunion.

  python3 scripts/generar_base_embotelladora.py [n]
"""
import csv
import subprocess
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "datos" / "embotelladora-200.csv"
DESTINO = RAIZ / "datos" / "embotelladora-200.xlsx"

AZUL = "1F3864"
NARANJA = "C55A11"
BLANCO = "FFFFFF"
cabecera = Font(bold=True, color=BLANCO, size=10)
centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
borde = Border(*[Side(style="thin", color="D9D9D9")] * 4)

ESCALA = [
    (5, "Su falla detiene la linea de envasado", "Llenadora, sopladora, capsuladora, lavadora"),
    (4, "Detiene un area o deja la linea sin insumo", "Jaraberia, tratamiento de agua, etiquetadora, empaque"),
    (3, "Degrada el proceso pero hay respaldo o by-pass", "Transporte, frio, aire comprimido, RILES"),
    (2, "Afecta calidad o eficiencia, no detiene", "Neumatica menor, lubricacion"),
    (1, "Sin impacto operacional directo", "Laboratorio, repuestos de taller"),
]


def main():
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 200
    subprocess.run(["node", "scripts/generar-embotelladora.mjs", str(n)], cwd=RAIZ, check=True,
                   capture_output=True)
    with open(ORIGEN, encoding="utf-8") as archivo:
        filas = list(csv.reader(archivo))
    encabezados, cuerpo = filas[0], filas[1:]

    libro = Workbook()
    libro.remove(libro.active)

    # -------------------------------------------------------------- Contexto
    h = libro.create_sheet("Contexto")
    h.column_dimensions["A"].width = 4
    h.column_dimensions["B"].width = 30
    h.column_dimensions["C"].width = 52
    h.column_dimensions["D"].width = 52
    h["B2"] = "MAESTRO DE REPUESTOS — PLANTA EMBOTELLADORA"
    h["B2"].font = Font(bold=True, size=16, color=AZUL)
    h["B3"] = (f"{len(cuerpo)} SKU de una linea de bebidas: desde el tratamiento de agua hasta el "
               "paletizado, mas servicios industriales y automatizacion.")
    h["B3"].font = Font(size=10, italic=True, color="808080")
    h["B5"] = ("Caso generico de la industria, construido para demostracion y prueba. No corresponde "
               "al maestro de ninguna planta real.")
    h["B5"].font = Font(size=10, color=NARANJA)

    areas = Counter(f[2] for f in cuerpo)
    h["B7"] = "AREAS CUBIERTAS"
    h["B7"].font = Font(bold=True, size=11, color=NARANJA)
    for i, texto in enumerate(["Area", "SKU", "Valor de inventario"], start=2):
        celda = h.cell(row=8, column=i, value=texto)
        celda.font = cabecera
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = centro
    idx_precio, idx_stock = encabezados.index("Precio unitario"), encabezados.index("Stock actual")
    valor_area = Counter()
    for f in cuerpo:
        valor_area[f[2]] += float(f[idx_precio]) * float(f[idx_stock])
    for i, (area, cuenta) in enumerate(sorted(areas.items(), key=lambda x: -valor_area[x[0]]), start=9):
        h.cell(row=i, column=2, value=area).font = Font(size=10)
        h.cell(row=i, column=3, value=cuenta).alignment = centro
        celda = h.cell(row=i, column=4, value=round(valor_area[area]))
        celda.number_format = '"US$"#,##0'
        for c in range(2, 5):
            h.cell(row=i, column=c).border = borde

    inicio = 9 + len(areas) + 2
    h.cell(row=inicio, column=2, value="ESCALA DE CRITICIDAD DEL EQUIPO").font = Font(
        bold=True, size=11, color=NARANJA)
    for i, texto in enumerate(["Nivel", "Significado", "Areas tipicas"], start=2):
        celda = h.cell(row=inicio + 1, column=i, value=texto)
        celda.font = cabecera
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = centro
    for i, (nivel, significado, areas_tipicas) in enumerate(ESCALA, start=inicio + 2):
        h.cell(row=i, column=2, value=nivel).alignment = centro
        h.cell(row=i, column=3, value=significado).font = Font(size=10)
        h.cell(row=i, column=4, value=areas_tipicas).font = Font(size=10, color="808080")
        for c in range(2, 5):
            h.cell(row=i, column=c).border = borde

    # ----------------------------------------------------------------- Datos
    d = libro.create_sheet("Datos")
    for i, texto in enumerate(encabezados, start=1):
        celda = d.cell(row=1, column=i, value=texto)
        celda.font = cabecera
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = centro
    numericas = {encabezados.index(c) + 1 for c in
                 ["Precio unitario", "Consumo anual", "Lead time (dias)",
                  "Criticidad equipo (1-5)", "Horas parada si falla", "N proveedores", "Stock actual"]}
    for f, registro in enumerate(cuerpo, start=2):
        for c, valor in enumerate(registro, start=1):
            celda = d.cell(row=f, column=c)
            if c in numericas:
                celda.value = float(valor) if "." in valor else int(valor)
                celda.number_format = "#,##0.00" if encabezados[c - 1] == "Precio unitario" else "#,##0"
            else:
                celda.value = valor
    anchos = [12, 46, 22, 30, 12, 9, 13, 15, 14, 15, 17, 17, 14, 13]
    for i, ancho in enumerate(anchos[:len(encabezados)], start=1):
        d.column_dimensions[get_column_letter(i)].width = ancho
    d.freeze_panes = "C2"
    d.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{len(cuerpo) + 1}"

    libro.save(DESTINO)
    print(f"OK  {DESTINO}  ({DESTINO.stat().st_size // 1024} KB, {len(cuerpo)} SKU, {len(areas)} areas)")


if __name__ == "__main__":
    main()
