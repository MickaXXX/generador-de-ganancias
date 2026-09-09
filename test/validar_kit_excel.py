#!/usr/bin/env python3
"""
Verificacion cruzada: las formulas de la planilla vendible deben dar exactamente
lo mismo que el motor JavaScript del sitio.

Regenera el kit, lo recalcula con LibreOffice y compara repuesto por repuesto el
indice de criticidad, la clase ABC, el stock de seguridad y el nivel objetivo S.

  python3 test/validar_kit_excel.py
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
KIT = RAIZ / "producto" / "Kit-Criticidad-Repuestos.xlsx"
FILAS = 60
TOLERANCIA = 1e-6

fallas = []


def comprobar(condicion, mensaje):
    print(f"{'  ok  ' if condicion else ' FALLA'} {mensaje}")
    if not condicion:
        fallas.append(mensaje)


def motor_js(fuente="demo", filas=FILAS):
    """Corre el mismo maestro por el motor del sitio."""
    generador = ("const filas = (await import('./scripts/generar-demo.mjs'))"
                 f".generarRepuestos({filas}, 7);" if fuente == "demo" else
                 "const filas = (await import('./scripts/generar-embotelladora.mjs'))"
                 f".generarEmbotelladora({filas});")
    guion = f"""
    import('./docs/core/criticidad.js').then(async (M) => {{
      {generador}
      const r = M.analizar(filas);
      console.log(JSON.stringify({{
        filas,
        pesos: r.pesos.map((p) => p.peso),
        items: Object.fromEntries(r.items.map((i) => [i.sku, {{
          C: i.indiceCriticidad, clase: i.clase, ss: i.stockSeguridad,
          S: i.nivelObjetivoS, liberable: i.capitalLiberable, riesgo: i.enRiesgo,
        }}])),
        resumen: r.resumen,
      }}));
    }});
    """
    salida = subprocess.run(["node", "-e", guion], cwd=RAIZ, capture_output=True, text=True, check=True)
    return json.loads(salida.stdout)


def recalcular(origen, destino_dir):
    """LibreOffice recalcula las formulas y guarda los valores en cache."""
    resultado = subprocess.run(
        ["soffice", "-env:UserInstallation=file:///tmp/lo-perfil-validacion",
         "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(destino_dir), str(origen)],
        capture_output=True, text=True, timeout=300)
    salida = destino_dir / origen.name
    if not salida.exists():
        raise RuntimeError(f"LibreOffice no pudo convertir: {resultado.stdout} {resultado.stderr}")
    return salida


def main():
    if not shutil.which("soffice"):
        print("Se necesita LibreOffice (soffice) para validar las formulas.")
        return 2

    subprocess.run([sys.executable, "scripts/generar_kit_excel.py", str(FILAS)],
                   cwd=RAIZ, check=True, capture_output=True)
    esperado = motor_js()

    with tempfile.TemporaryDirectory() as tmp:
        recalculado = recalcular(KIT, Path(tmp))
        libro = load_workbook(recalculado, data_only=True)
        crit = libro["Criticidad"]
        resumen = libro["Resumen"]

        print("\n== Pesos por entropia ==")
        for i, peso_js in enumerate(esperado["pesos"]):
            peso_excel = crit[f"BE{2 + i}"].value
            comprobar(peso_excel is not None and abs(peso_excel - peso_js) < TOLERANCIA,
                      f"criterio {i + 1}: Excel {peso_excel!r} vs motor {peso_js:.9f}")

        print("\n== Repuesto por repuesto ==")
        diferencias = {"C": 0, "clase": 0, "ss": 0, "S": 0, "riesgo": 0}
        comparados = 0
        for fila in range(2, 2 + FILAS):
            sku = crit[f"A{fila}"].value
            if not sku:
                continue
            ref = esperado["items"].get(sku)
            if ref is None:
                fallas.append(f"SKU {sku} no existe en el motor")
                continue
            comparados += 1
            if abs((crit[f"AG{fila}"].value or 0) - ref["C"]) > 1e-9:
                diferencias["C"] += 1
            if crit[f"AJ{fila}"].value != ref["clase"]:
                diferencias["clase"] += 1
            if (crit[f"AO{fila}"].value or 0) != ref["ss"]:
                diferencias["ss"] += 1
            if (crit[f"AP{fila}"].value or 0) != ref["S"]:
                diferencias["S"] += 1
            if (crit[f"AT{fila}"].value == "SI") != ref["riesgo"]:
                diferencias["riesgo"] += 1

        comprobar(comparados == FILAS, f"se compararon los {FILAS} repuestos ({comparados})")
        for campo, cuenta in diferencias.items():
            comprobar(cuenta == 0, f"{campo}: {cuenta} diferencias de {comparados}")

        print("\n== Totales del resumen ==")
        n_excel = resumen["C6"].value
        a_excel = resumen["C7"].value
        liberable_excel = resumen["C13"].value or 0
        riesgo_excel = resumen["C17"].value
        pct_excel = resumen["C14"].value
        capital_excel = resumen["C12"].value or 0

        comprobar(n_excel == esperado["resumen"]["n"], f"repuestos: {n_excel} vs {esperado['resumen']['n']}")
        comprobar(a_excel == esperado["resumen"]["porClase"]["A"],
                  f"clase A: {a_excel} vs {esperado['resumen']['porClase']['A']}")
        comprobar(abs(liberable_excel - esperado["resumen"]["capitalLiberable"]) < 0.01,
                  f"capital liberable: {liberable_excel:,.2f} vs {esperado['resumen']['capitalLiberable']:,.2f}")
        comprobar(abs(capital_excel - esperado["resumen"]["capitalActual"]) < 0.01,
                  f"capital inmovilizado: {capital_excel:,.2f} vs {esperado['resumen']['capitalActual']:,.2f}")
        comprobar(riesgo_excel == esperado["resumen"]["enRiesgo"],
                  f"bajo cobertura: {riesgo_excel} vs {esperado['resumen']['enRiesgo']}")
        comprobar(pct_excel not in (None, 0) and 0 < pct_excel < 1,
                  f"% de capital que sobra calculado: {pct_excel}")

        print("\n== Estructura del archivo ==")
        for hoja in ("Resumen", "Instrucciones", "Datos", "Criticidad", "Parametros"):
            comprobar(hoja in libro.sheetnames, f"existe la hoja «{hoja}»")
        comprobar(libro.sheetnames[0] == "Resumen", "el archivo abre en el Resumen")

        # Una fila vacia no debe ensuciar la planilla con errores.
        vacias = [crit[f"AG{f}"].value for f in range(2 + FILAS, 2 + FILAS + 5)]
        comprobar(all(v in (None, "") for v in vacias), f"las filas sin datos quedan limpias: {vacias}")

    validar_embotelladora()

    print(f"\n{'TODO OK' if not fallas else f'{len(fallas)} FALLAS'}\n")
    return 0 if not fallas else 1


def validar_embotelladora(n=200):
    """
    El caso duro: un maestro con varios ordenes de magnitud de diferencia en
    precio y consumo, que es donde el metodo se rompia. Se pegan esas filas en
    la hoja Datos del kit y se comprueba que Excel siga coincidiendo con el motor.
    """
    from openpyxl import load_workbook as cargar

    print(f"\n== Maestro de embotelladora ({n} SKU, rango dinamico amplio) ==")
    esperado = motor_js("embotelladora", n)
    claves = ["sku", "descripcion", "categoria", "precio", "consumoAnual", "leadTime",
              "criticidadEquipo", "horasParada", "proveedores", "stockActual"]

    libro = cargar(KIT)
    datos = libro["Datos"]
    for fila in range(2, 2 + 300):          # se limpian las filas de ejemplo
        for col in range(1, 11):
            datos.cell(row=fila, column=col).value = None
    for i, registro in enumerate(esperado["filas"], start=2):
        for c, clave in enumerate(claves, start=1):
            datos.cell(row=i, column=c).value = registro.get(clave)

    with tempfile.TemporaryDirectory() as tmp:
        origen = Path(tmp) / "kit-embotelladora.xlsx"
        libro.save(origen)
        salida = Path(tmp) / "recalc"
        salida.mkdir()
        recalculado = recalcular(origen, salida)
        crit = cargar(recalculado, data_only=True)["Criticidad"]

        diferencias = {"C": 0, "clase": 0, "ss": 0, "S": 0}
        comparados = 0
        for fila in range(2, 2 + n):
            sku = crit[f"A{fila}"].value
            ref = esperado["items"].get(sku)
            if ref is None:
                continue
            comparados += 1
            if abs((crit[f"AG{fila}"].value or 0) - ref["C"]) > 1e-9:
                diferencias["C"] += 1
            if crit[f"AJ{fila}"].value != ref["clase"]:
                diferencias["clase"] += 1
            if (crit[f"AO{fila}"].value or 0) != ref["ss"]:
                diferencias["ss"] += 1
            if (crit[f"AP{fila}"].value or 0) != ref["S"]:
                diferencias["S"] += 1

        comprobar(comparados == n, f"se compararon los {n} repuestos ({comparados})")
        for campo, cuenta in diferencias.items():
            comprobar(cuenta == 0, f"{campo}: {cuenta} diferencias de {comparados}")
        clases = [crit[f"AJ{f}"].value for f in range(2, 2 + n)]
        comprobar(clases.count("A") <= n * 0.21,
                  f"la clase A se mantiene acotada en Excel: {clases.count('A')} de {n}")


if __name__ == "__main__":
    sys.exit(main())
